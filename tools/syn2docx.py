# NOTE: superseded for the synopsis. docs/sys.docx is now the master -
# its typography was changed by hand in Word, and running this script would
# overwrite that. Kept for reference and for regenerating other documents.

"""Build the synopsis .docx from docs/synopsis.html.

Generated from the same source as the PDF on purpose: a hand-maintained Word
copy drifts from the PDF within one edit, and then the two versions disagree in
front of the department. Run this after editing synopsis.html.

    python syn2docx.py <synopsis.html> <out.docx>
"""

from __future__ import annotations

import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

NL = chr(10)  # a hard line break inside a docx run
BLUE = RGBColor(0x2E, 0x74, 0xB5)
HDR_FILL = "DEEAF6"
BODY_PT = 9.7


# --------------------------------------------------------------- html parsing


class Node:
    """A node keeps its text and children in ONE ordered list.

    Storing them separately loses position: `<p>A <b>B</b> C</p>` would emit
    "A  C" then "B", which silently scrambles every sentence containing an
    inline tag or a <br>.
    """

    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.parts = []  # str | Node, in document order

    @property
    def children(self):
        return [p for p in self.parts if isinstance(p, Node)]

    @property
    def text(self):
        return "".join(p for p in self.parts if isinstance(p, str))

    @property
    def cls(self):
        return self.attrs.get("class", "")


class Tree(HTMLParser):
    """Minimal DOM builder - the synopsis markup is ours and stays simple."""

    VOID = {"br", "hr", "img", "meta", "link"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs)
        self.stack[-1].parts.append(n)
        if tag not in self.VOID:
            self.stack.append(n)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].parts.append(Node(tag, attrs))

    def handle_endtag(self, tag):
        if tag in self.VOID:
            return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].parts.append(data)


def find(node, tag=None, cls=None):
    out = []
    for c in node.children:
        if (tag is None or c.tag == tag) and (cls is None or cls in c.cls.split()):
            out.append(c)
        out.extend(find(c, tag, cls))
    return out


def runs_of(node):
    """Flatten a node into ordered (text, bold, italic, code) runs."""
    out = []

    def walk(n, b, i, c):
        for part in n.parts:
            if isinstance(part, str):
                out.append((part, b, i, c))
            elif part.tag == "br":
                out.append((NL, b, i, c))
            elif part.tag == "hr":
                continue
            else:
                walk(
                    part,
                    b or part.tag in ("b", "strong"),
                    i or part.tag in ("i", "em"),
                    c or part.tag == "code",
                )

    walk(node, False, False, False)

    merged = []
    for t, b, i, c in out:
        t = t if t == NL else re.sub(r"\s+", " ", t)
        if not t:
            continue
        if merged and merged[-1][0] != NL and t != NL and merged[-1][1:] == (b, i, c):
            merged[-1] = (merged[-1][0] + t, b, i, c)
        else:
            merged.append((t, b, i, c))

    while merged and merged[0][0].strip() == "" and merged[0][0] != NL:
        merged.pop(0)
    while merged and merged[-1][0].strip() == "" and merged[-1][0] != NL:
        merged.pop()
    if merged:
        merged[0] = (merged[0][0].lstrip(), *merged[0][1:])
        merged[-1] = (merged[-1][0].rstrip(), *merged[-1][1:])
    return [m for m in merged if m[0]]


# ------------------------------------------------------------- docx utilities


def style_run(r, *, size=BODY_PT, bold=False, italic=False, color=None, mono=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Consolas" if mono else "Calibri"
    if color is not None:
        r.font.color.rgb = color
    return r


def para(doc, text="", *, align=None, size=BODY_PT, bold=False, italic=False,
         color=None, before=0, after=3, line=1.15):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if text:
        style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def emit_runs(p, node, size=BODY_PT):
    for t, b, i, c in runs_of(node):
        if t == "\n":
            p.add_run().add_break()
            continue
        style_run(p.add_run(t), size=size * (0.95 if c else 1), bold=b, italic=i, mono=c)


def shade(cell, fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def page_border(section):
    """A4 cover box. python-docx has no API for this, so it is raw XML."""
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "12")       # 1.5pt
        e.set(qn("w:space"), "24")    # distance from page edge
        e.set(qn("w:color"), "000000")
        borders.append(e)
    sectPr.append(borders)


def clear_page_border(section):
    """`add_section` clones the previous sectPr, border included."""
    sectPr = section._sectPr
    for b in sectPr.findall(qn("w:pgBorders")):
        sectPr.remove(b)


def set_margins(section, top, right, bottom, left):
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin, section.right_margin = Mm(top), Mm(right)
    section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)


# ------------------------------------------------------------------ the build


def build(html_path: Path, out_path: Path) -> None:
    tree = Tree()
    tree.feed(html_path.read_text(encoding="utf-8"))
    root = tree.root

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(BODY_PT)

    # ---------------------------------------------------------- cover page
    sec = doc.sections[0]
    set_margins(sec, 26, 24, 20, 24)
    page_border(sec)

    cover = find(root, cls="cover")[0]
    C = WD_ALIGN_PARAGRAPH.CENTER

    def cov(cls):
        return find(cover, cls=cls)

    mp = cov("mp")[0]
    p = para(doc, align=C, after=2, before=10)
    for t, *_ in runs_of(mp):
        if t == "\n":
            p.add_run().add_break()
        else:
            style_run(p.add_run(t), size=13, bold=True)

    para(doc, cov("title")[0].text.strip(), align=C, size=21, color=BLUE,
         before=10, after=6, line=1.1)

    # the blue rule under the title
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(16)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pbdr.append(bottom)
    rule._p.get_or_add_pPr().append(pbdr)

    para(doc, "by", align=C, italic=True, after=10)

    names = cov("nm")
    entries = cov("en")
    for nm, en in zip(names, entries):
        para(doc, nm.text.strip(), align=C, size=11, bold=True, before=7, after=0)
        para(doc, en.text.strip(), align=C, size=11, bold=True, after=0)

    g = para(doc, align=C, before=18, after=0, line=1.6)
    for t, *_ in runs_of(cov("guide")[0]):
        if t == "\n":
            g.add_run().add_break()
        else:
            style_run(g.add_run(t), bold=True)

    para(doc, cov("partial")[0].text.strip(), align=C, italic=True, size=9.6,
         before=18, after=2)
    brs = cov("br")
    para(doc, brs[0].text.strip(), align=C, size=10.5, after=2)
    para(doc, cov("deg")[0].text.strip(), align=C, size=11, bold=True, after=2)
    para(doc, "in", align=C, size=10.5, italic=True, after=2)
    para(doc, brs[-1].text.strip(), align=C, size=10.5, after=4)

    logo = html_path.parent / find(cover, tag="img")[0].attrs.get("src", "")
    if logo.exists():
        lp = doc.add_paragraph()
        lp.paragraph_format.alignment = C
        lp.paragraph_format.space_before = Pt(12)
        lp.paragraph_format.space_after = Pt(8)
        lp.add_run().add_picture(str(logo), width=Inches(0.92))

    para(doc, cov("uni")[0].text.strip(), align=C, size=11.5, bold=True, after=4)
    para(doc, cov("sch")[0].text.strip(), align=C, size=10.5, bold=True, after=5)
    ap = para(doc, align=C, after=0, line=1.4)
    for t, *_ in runs_of(cov("addr")[0]):
        if t == "\n":
            ap.add_run().add_break()
        else:
            style_run(ap.add_run(t), size=10.5, bold=True)

    # ------------------------------------------------------- content pages
    content = doc.add_section(WD_SECTION.NEW_PAGE)
    clear_page_border(content)
    set_margins(content, 15, 16, 13, 16)

    page = find(root, cls="page")[0]
    for node in page.children:
        if node.tag == "h2":
            para(doc, node.text.strip(), size=10.6, bold=True, color=BLUE,
                 before=9, after=3)

        elif node.tag == "p":
            p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)
            emit_runs(p, node)

        elif node.tag in ("ul", "ol"):
            refs = "refs" in node.cls
            size = 8.5 if refs else BODY_PT
            items = [c for c in node.children if c.tag == "li"]
            for n, li in enumerate(items, start=1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(1.5)
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf.left_indent = Inches(0.32)
                pf.first_line_indent = Inches(-0.32)
                marker = "•	" if node.tag == "ul" else f"{n}.	"
                style_run(p.add_run(marker), size=size)
                emit_runs(p, li, size=size)

        elif node.tag == "table":
            rows = [r for r in find(node, tag="tr")]
            cells0 = [c for c in rows[0].children if c.tag in ("th", "td")]
            t = doc.add_table(rows=0, cols=len(cells0))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r in rows:
                cs = [c for c in r.children if c.tag in ("th", "td")]
                row = t.add_row()
                for cell, src in zip(row.cells, cs):
                    cell.paragraphs[0].text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    emit_runs(p, src, size=8.7)
                    if src.tag == "th":
                        shade(cell, HDR_FILL)
                        for r_ in p.runs:
                            r_.font.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    doc.save(out_path)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
